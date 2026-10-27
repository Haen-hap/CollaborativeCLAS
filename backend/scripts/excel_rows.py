"""Lee la hoja Raw_Sessions_Data del Excel del cliente y devuelve diccionarios listos para SessionIn."""
from openpyxl import load_workbook


def _bool(v) -> bool:
    return v is True or str(v).strip().upper() == "TRUE"


def read_excel(path: str) -> list[dict]:
    ws = load_workbook(path, read_only=True, data_only=True)["Raw_Sessions_Data"]
    rows = ws.iter_rows(values_only=True)
    header = [str(h) for h in next(rows)]
    out = []
    for r in rows:
        if not r or r[0] is None:
            continue
        d = dict(zip(header, r))
        d["assisted_mode"] = _bool(d["assisted_mode"])
        d["reward_issued"] = _bool(d["reward_issued"])
        d["origin"] = "UNKNOWN" if d["kiosk_id"] == "DIRECT_WEB" else "KIOSK"
        out.append(d)
    return out
